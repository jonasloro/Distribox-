from __future__ import annotations

import json
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, FastAPI, File, HTTPException, UploadFile
from openpyxl import load_workbook

from supabase_module import pg_connect

BASE_DIR = Path(__file__).resolve().parent
router = APIRouter(prefix="/api/production", tags=["Produção"])

SECTORS = ["Triagem", "Qualidade", "Processamento", "Etiquetagem", "Estocagem", "Expedição"]

QTY_HEADERS = {"quantidade", "quantidade feita", "quantidade total", "quantidade de volumes"}
RESP_HEADERS = {"responsavel"}
DATA_HEADERS = {"data"}
INICIO_HEADERS = {"inicio"}
TERMINO_HEADERS = {"termino"}
SKIP_HEADERS = {"pausas", "inicio 2", "termino 2", "duracao 2", "duracao", "total"}


def iso_now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def init_production_db() -> None:
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                cur.execute(
                    """
                    CREATE TABLE IF NOT EXISTS production_entries (
                        id SERIAL PRIMARY KEY,
                        sector TEXT NOT NULL,
                        responsavel TEXT NOT NULL,
                        campos TEXT,
                        quantidade INTEGER NOT NULL DEFAULT 0,
                        data TEXT,
                        inicio TEXT,
                        termino TEXT,
                        pausas_seconds INTEGER NOT NULL DEFAULT 0,
                        pause_started_at TEXT,
                        duracao_seconds INTEGER,
                        status TEXT NOT NULL DEFAULT 'EM_ANDAMENTO',
                        source TEXT NOT NULL DEFAULT 'APP',
                        created_at TEXT NOT NULL
                    );
                    CREATE INDEX IF NOT EXISTS idx_production_sector ON production_entries(sector, status);
                    CREATE TABLE IF NOT EXISTS production_feedback (
                        id SERIAL PRIMARY KEY,
                        entry_id INTEGER NOT NULL,
                        sector TEXT NOT NULL,
                        teve_dificuldade INTEGER NOT NULL,
                        comentario TEXT,
                        created_at TEXT NOT NULL
                    );
                    """
                )
            con.commit()
    except RuntimeError as e:
        print(f"[aviso] Produção: {e}")


def _row(r: dict[str, Any]) -> dict[str, Any]:
    d = dict(r)
    d["campos"] = json.loads(d["campos"]) if d.get("campos") else {}
    return d


@router.post("/entries")
def start_entry(payload: dict[str, Any]) -> dict[str, Any]:
    sector = str(payload.get("sector") or "").strip()
    responsavel = str(payload.get("responsavel") or "").strip()
    if sector not in SECTORS:
        raise HTTPException(400, "Setor inválido.")
    if not responsavel:
        raise HTTPException(400, "Informe o responsável.")
    now = iso_now()
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                cur.execute(
                    """INSERT INTO production_entries(sector,responsavel,campos,quantidade,data,inicio,status,created_at)
                       VALUES(%s,%s,%s,%s,%s,%s,'EM_ANDAMENTO',%s) RETURNING *""",
                    (sector, responsavel, json.dumps(payload.get("campos") or {}, ensure_ascii=False),
                     int(payload.get("quantidade") or 0), date.today().isoformat(), now, now),
                )
                entry = cur.fetchone()
            con.commit()
    except RuntimeError as e:
        raise HTTPException(503, str(e))
    return _row(entry)


def _get_entry(cur, entry_id: int) -> dict[str, Any]:
    cur.execute("SELECT * FROM production_entries WHERE id=%s", (entry_id,))
    entry = cur.fetchone()
    if not entry:
        raise HTTPException(404, "Registro não encontrado.")
    return entry


@router.post("/entries/{entry_id}/pause")
def pause_entry(entry_id: int) -> dict[str, Any]:
    with pg_connect() as con:
        with con.cursor() as cur:
            entry = _get_entry(cur, entry_id)
            if entry["status"] != "EM_ANDAMENTO":
                raise HTTPException(400, "Só é possível pausar uma tarefa em andamento.")
            cur.execute(
                "UPDATE production_entries SET status='PAUSADO', pause_started_at=%s WHERE id=%s RETURNING *",
                (iso_now(), entry_id),
            )
            entry = cur.fetchone()
        con.commit()
    return _row(entry)


@router.post("/entries/{entry_id}/resume")
def resume_entry(entry_id: int) -> dict[str, Any]:
    with pg_connect() as con:
        with con.cursor() as cur:
            entry = _get_entry(cur, entry_id)
            if entry["status"] != "PAUSADO":
                raise HTTPException(400, "Essa tarefa não está pausada.")
            elapsed = (datetime.fromisoformat(iso_now()) - datetime.fromisoformat(entry["pause_started_at"])).total_seconds()
            cur.execute(
                """UPDATE production_entries SET status='EM_ANDAMENTO', pausas_seconds=pausas_seconds+%s,
                   pause_started_at=NULL WHERE id=%s RETURNING *""",
                (int(elapsed), entry_id),
            )
            entry = cur.fetchone()
        con.commit()
    return _row(entry)


@router.post("/entries/{entry_id}/complete")
def complete_entry(entry_id: int, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = payload or {}
    with pg_connect() as con:
        with con.cursor() as cur:
            entry = _get_entry(cur, entry_id)
            if entry["status"] == "CONCLUIDO":
                raise HTTPException(400, "Essa tarefa já foi concluída.")
            pausas = entry["pausas_seconds"]
            if entry["status"] == "PAUSADO" and entry["pause_started_at"]:
                pausas += int((datetime.fromisoformat(iso_now()) - datetime.fromisoformat(entry["pause_started_at"])).total_seconds())
            termino = iso_now()
            total = int((datetime.fromisoformat(termino) - datetime.fromisoformat(entry["inicio"])).total_seconds())
            duracao = max(0, total - pausas)
            quantidade = payload.get("quantidade")
            cur.execute(
                """UPDATE production_entries SET status='CONCLUIDO', termino=%s, pausas_seconds=%s,
                   pause_started_at=NULL, duracao_seconds=%s, quantidade=COALESCE(%s,quantidade)
                   WHERE id=%s RETURNING *""",
                (termino, pausas, duracao, quantidade, entry_id),
            )
            entry = cur.fetchone()
        con.commit()
    return _row(entry)


@router.get("/entries")
def list_entries(sector: Optional[str] = None, status: Optional[str] = None, limit: int = 20) -> list[dict[str, Any]]:
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                query = "SELECT * FROM production_entries WHERE 1=1"
                params: list[Any] = []
                if sector:
                    query += " AND sector=%s"
                    params.append(sector)
                if status:
                    query += " AND status=%s"
                    params.append(status)
                query += " ORDER BY id DESC LIMIT %s"
                params.append(max(1, min(limit, 200)))
                cur.execute(query, params)
                rows = cur.fetchall()
    except RuntimeError:
        return []
    return [_row(r) for r in rows]


@router.get("/by-person")
def production_by_person(sector: Optional[str] = None) -> dict[str, Any]:
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                query = "SELECT sector, responsavel, quantidade FROM production_entries WHERE status='CONCLUIDO'"
                params: list[Any] = []
                if sector:
                    query += " AND sector=%s"
                    params.append(sector)
                cur.execute(query, params)
                rows = cur.fetchall()
    except RuntimeError:
        rows = []
    sectors: dict[str, dict[str, dict[str, Any]]] = {}
    for r in rows:
        agg = sectors.setdefault(r["sector"], {})
        entry = agg.setdefault(r["responsavel"], {"name": r["responsavel"], "tasks": 0, "qty": 0})
        entry["tasks"] += 1
        entry["qty"] += int(r["quantidade"] or 0)
    result = {s: sorted(people.values(), key=lambda x: -x["qty"]) for s, people in sectors.items()}
    totals = {s: {"tasks": sum(p["tasks"] for p in people), "qty": sum(p["qty"] for p in people)} for s, people in result.items()}
    return {"sectors": result, "totals": totals}


@router.post("/entries/{entry_id}/feedback")
def submit_feedback(entry_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    with pg_connect() as con:
        with con.cursor() as cur:
            entry = _get_entry(cur, entry_id)
            cur.execute(
                """INSERT INTO production_feedback(entry_id,sector,teve_dificuldade,comentario,created_at)
                   VALUES(%s,%s,%s,%s,%s)""",
                (entry_id, entry["sector"], 1 if payload.get("teve_dificuldade") else 0,
                 (payload.get("comentario") or "").strip() or None, iso_now()),
            )
        con.commit()
    return {"ok": True}


def _cell_to_text(value: Any) -> Any:
    if isinstance(value, (datetime, date, dtime)):
        return value.isoformat()
    return value


@router.post("/import")
async def import_sheet(sector: str, file: UploadFile = File(...)) -> dict[str, Any]:
    if sector not in SECTORS:
        raise HTTPException(400, "Setor inválido.")
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Envie um arquivo Excel .xlsx ou .xlsm.")
    content = await file.read()
    tmp = BASE_DIR / "data" / f"import_producao_{sector}.xlsx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    tmp.write_bytes(content)
    wb = load_workbook(tmp, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    raw_headers = next(rows_iter, None)
    if not raw_headers:
        raise HTTPException(400, "A planilha está vazia.")
    headers = [str(h).strip() if h else "" for h in raw_headers]
    norm = [h.lower() for h in headers]

    imported = 0
    skipped = 0
    batch: list[tuple] = []
    for row in rows_iter:
        if not any(v is not None for v in row):
            continue
        values = dict(zip(norm, row))
        headers_map = dict(zip(norm, headers))
        responsavel = values.get("responsavel")
        if not responsavel:
            skipped += 1
            continue
        qty = 0
        for h in QTY_HEADERS:
            if values.get(h) not in (None, ""):
                try:
                    qty = int(values.get(h))
                except (TypeError, ValueError):
                    qty = 0
                break
        data_val = values.get("data")
        data_iso = data_val.date().isoformat() if isinstance(data_val, (datetime, date)) else None
        inicio_val = values.get("inicio")
        termino_val = values.get("termino")
        inicio_iso = termino_iso = None
        duracao = None
        if isinstance(data_val, (datetime, date)) and isinstance(inicio_val, (dtime, datetime)):
            base_date = data_val.date() if isinstance(data_val, datetime) else data_val
            t = inicio_val.time() if isinstance(inicio_val, datetime) else inicio_val
            inicio_iso = datetime.combine(base_date, t).isoformat()
            if isinstance(termino_val, (dtime, datetime)):
                t2 = termino_val.time() if isinstance(termino_val, datetime) else termino_val
                termino_iso = datetime.combine(base_date, t2).isoformat()
                duracao = int((datetime.fromisoformat(termino_iso) - datetime.fromisoformat(inicio_iso)).total_seconds())
                if duracao < 0:
                    duracao = None
        campos = {
            headers_map[k]: _cell_to_text(v)
            for k, v in values.items()
            if k not in RESP_HEADERS | QTY_HEADERS | DATA_HEADERS | INICIO_HEADERS | TERMINO_HEADERS | SKIP_HEADERS
            and v not in (None, "")
        }
        batch.append((
            sector, str(responsavel).strip(), json.dumps(campos, ensure_ascii=False), qty,
            data_iso, inicio_iso, termino_iso, duracao, iso_now(),
        ))
        imported += 1

    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                if batch:
                    cur.executemany(
                        """INSERT INTO production_entries(sector,responsavel,campos,quantidade,data,inicio,termino,duracao_seconds,status,source,created_at)
                           VALUES(%s,%s,%s,%s,%s,%s,%s,%s,'CONCLUIDO','IMPORTACAO',%s)""",
                        batch,
                    )
            con.commit()
    except RuntimeError as e:
        raise HTTPException(503, str(e))
    return {"sector": sector, "imported": imported, "skipped": skipped}


def register_production_routes(app: FastAPI) -> None:
    app.include_router(router)
