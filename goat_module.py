from __future__ import annotations

import json
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, FastAPI, File, HTTPException, UploadFile
from openpyxl import load_workbook

from supabase_module import pg_connect

BASE_DIR = Path(__file__).resolve().parent
router = APIRouter(prefix="/api/goat", tags=["GOAT"])

# Abas conhecidas do relatório "Tudo das operações" do GOAT. Se o GOAT mandar uma
# aba com outro nome, ela ainda é importada, só não recebe um rótulo amigável aqui.
SHEET_LABELS = {
    "Rejeitados": "Peças rejeitadas por lote/fornecedor",
    "Retrabalho": "Peças em retrabalho",
    "Concluídos": "Lotes concluídos",
    "Separado para foto": "Fila de amostras para foto",
    "Devoluções": "Devoluções ao fornecedor",
    "Crédito com fornecedor": "Créditos em aberto com fornecedor",
    "Acerto por fornecedor": "Resumo de acerto por fornecedor",
    "Abatimento no boleto": "Abatimentos no boleto",
    "Perdas": "Perdas registradas",
    "Desempenho de fornecedores": "Score de desempenho por fornecedor",
    "Entregas atrasadas": "Entregas em atraso",
    "Compra x recebido": "Divergência compra x recebido",
    "Estocagem x esperado": "Divergência de contagem física",
}


def init_goat_db() -> None:
    """Cria as tabelas no Supabase se ainda não existirem. Se o Supabase não
    estiver configurado (SUPABASE_DATABASE_URL ausente), só avisa e segue —
    mesmo comportamento do resto do app nessa situação."""
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                cur.execute(
                    """
                    CREATE TABLE IF NOT EXISTS goat_imports (
                        id SERIAL PRIMARY KEY,
                        filename TEXT,
                        sheets_count INTEGER,
                        rows_count INTEGER,
                        imported_at TEXT NOT NULL
                    );
                    CREATE TABLE IF NOT EXISTS goat_records (
                        id SERIAL PRIMARY KEY,
                        sheet TEXT NOT NULL,
                        data TEXT NOT NULL,
                        imported_at TEXT NOT NULL
                    );
                    CREATE INDEX IF NOT EXISTS idx_goat_records_sheet ON goat_records(sheet);
                    """
                )
            con.commit()
    except RuntimeError as e:
        print(f"[aviso] GOAT: {e}")


def _cell(v: Any) -> Any:
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    return v


def _num(v: Any) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


@router.post("/import")
async def import_report(file: UploadFile = File(...)) -> dict[str, Any]:
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Envie um arquivo Excel .xlsx ou .xlsm.")
    content = await file.read()
    tmp = BASE_DIR / "data" / "import_goat.xlsx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    tmp.write_bytes(content)
    wb = load_workbook(tmp, read_only=True, data_only=True)

    now = datetime.now().replace(microsecond=0).isoformat()
    total_rows = 0
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                # Cada importação substitui o snapshot anterior — é um relatório
                # fotográfico, não um histórico incremental (igual o SGO fazia).
                cur.execute("DELETE FROM goat_records")
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows_iter = ws.iter_rows(values_only=True)
                    headers = next(rows_iter, None)
                    if not headers:
                        continue
                    headers = [str(h).strip() if h else f"col{i}" for i, h in enumerate(headers)]
                    batch: list[tuple[str, str, str]] = []
                    for row in rows_iter:
                        if not any(v is not None for v in row):
                            continue
                        record = {headers[i]: _cell(row[i]) for i in range(min(len(headers), len(row)))}
                        batch.append((sheet_name, json.dumps(record, ensure_ascii=False), now))
                    if batch:
                        cur.executemany(
                            "INSERT INTO goat_records(sheet,data,imported_at) VALUES(%s,%s,%s)", batch
                        )
                        total_rows += len(batch)
                cur.execute(
                    "INSERT INTO goat_imports(filename,sheets_count,rows_count,imported_at) VALUES(%s,%s,%s,%s)",
                    (file.filename, len(wb.sheetnames), total_rows, now),
                )
            con.commit()
    except RuntimeError as e:
        raise HTTPException(503, str(e))
    return {"sheets": len(wb.sheetnames), "rows": total_rows, "imported_at": now}


def _sheet_rows(cur, sheet: str) -> list[dict[str, Any]]:
    cur.execute("SELECT data FROM goat_records WHERE sheet=%s", (sheet,))
    return [json.loads(r["data"]) for r in cur.fetchall()]


@router.get("/status")
def status() -> dict[str, Any]:
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                cur.execute("SELECT * FROM goat_imports ORDER BY id DESC LIMIT 1")
                last = cur.fetchone()
                cur.execute("SELECT sheet, COUNT(*) c FROM goat_records GROUP BY sheet")
                sheets = cur.fetchall()
    except RuntimeError:
        return {"last_import": None, "sheets": []}
    return {
        "last_import": dict(last) if last else None,
        "sheets": [{"sheet": s["sheet"], "label": SHEET_LABELS.get(s["sheet"], s["sheet"]), "rows": s["c"]} for s in sheets],
    }


@router.get("/sheet/{sheet}")
def sheet_rows(sheet: str, limit: int = 100) -> dict[str, Any]:
    with pg_connect() as con:
        with con.cursor() as cur:
            rows = _sheet_rows(cur, sheet)
    return {"sheet": sheet, "label": SHEET_LABELS.get(sheet, sheet), "total": len(rows), "rows": rows[: max(1, min(limit, 500))]}


@router.get("/summary")
def summary() -> dict[str, Any]:
    try:
        with pg_connect() as con:
            with con.cursor() as cur:
                rejeitados = _sheet_rows(cur, "Rejeitados")
                perdas = _sheet_rows(cur, "Perdas")
                desempenho = _sheet_rows(cur, "Desempenho de fornecedores")
                atrasadas = _sheet_rows(cur, "Entregas atrasadas")
                divergencia_compra = _sheet_rows(cur, "Compra x recebido")
                divergencia_estoque = _sheet_rows(cur, "Estocagem x esperado")
    except RuntimeError:
        rejeitados = perdas = desempenho = atrasadas = divergencia_compra = divergencia_estoque = []

    por_fornecedor: dict[str, dict[str, Any]] = {}
    for r in rejeitados:
        forn = r.get("Fornecedor") or "—"
        entry = por_fornecedor.setdefault(forn, {"name": forn, "pecas": 0, "valor": 0.0})
        entry["pecas"] += int(_num(r.get("Peças")))
        entry["valor"] += _num(r.get("Valor"))
    top_rejeitados = sorted(por_fornecedor.values(), key=lambda x: -x["pecas"])[:8]

    por_responsavel: dict[str, dict[str, Any]] = {}
    for r in perdas:
        quem = r.get("Quem") or "—"
        entry = por_responsavel.setdefault(quem, {"name": quem, "pecas": 0})
        entry["pecas"] += int(_num(r.get("Peças")))
    top_perdas = sorted(por_responsavel.values(), key=lambda x: -x["pecas"])[:8]

    top_desempenho = sorted(desempenho, key=lambda x: -_num(x.get("Atrasadas hoje")))[:8]

    return {
        "totals": {
            "rejeitados_pecas": sum(int(_num(r.get("Peças"))) for r in rejeitados),
            "rejeitados_valor": round(sum(_num(r.get("Valor")) for r in rejeitados), 2),
            "perdas_pecas": sum(int(_num(r.get("Peças"))) for r in perdas),
            "entregas_atrasadas": len(atrasadas),
            "divergencias_compra": sum(1 for r in divergencia_compra if _num(r.get("Diferença")) != 0),
            "divergencias_estoque": sum(1 for r in divergencia_estoque if _num(r.get("Diferença")) != 0),
        },
        "top_rejeitados_fornecedor": top_rejeitados,
        "top_perdas_responsavel": top_perdas,
        "fornecedores_criticos": [
            {
                "name": d.get("Fornecedor"),
                "atrasadas_hoje": int(_num(d.get("Atrasadas hoje"))),
                "acerto_pct": _num(d.get("Acerto da data (%)")),
            }
            for d in top_desempenho
        ],
    }


def register_goat_routes(app: FastAPI) -> None:
    app.include_router(router)
